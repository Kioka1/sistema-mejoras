"""

Servidor (backend) del Sistema de Registro y Seguimiento de Mejoras

e Incidencias.



Esta es la parte "Servidor" de la arquitectura Cliente-Servidor: recibe

peticiones del frontend (React), consulta la base de datos, y devuelve

la información en formato JSON. El frontend nunca toca la base de datos

directamente, siempre pasa por esta API.



Cómo correrlo:

    cd backend

    ./venv/bin/uvicorn main:app --reload



La documentación interactiva de la API queda disponible automáticamente

en http://localhost:8000/docs

"""



from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form

from fastapi.middleware.cors import CORSMiddleware

from fastapi.responses import StreamingResponse, FileResponse

from sqlalchemy.orm import Session, joinedload

from sqlalchemy import or_, func

from typing import Optional, List

import os

import re

import uuid

import time

import threading

from contextlib import asynccontextmanager

from datetime import datetime, date, timedelta

from io import BytesIO

from openpyxl import Workbook, load_workbook

from openpyxl.styles import Font, PatternFill

from google.oauth2 import id_token as google_id_token

from google.auth.transport import requests as google_requests



from database import (

    Base, engine, get_db, SessionLocal, Registro, AccionMejora, HistorialAccion, Solicitud, HistorialSolicitud,

    Usuario, EvidenciaArchivo, Recomendacion, Area, PlanAccionItem, Notificacion,

    parse_areas, asignar_areas, obtener_o_crear_area,

)

from auth import obtener_usuario_actual, obtener_usuario_para_descarga, requerir_admin, requerir_admin_descarga, verificar_password, crear_token, hashear_password

from notificaciones import enviar_correo

from schemas import (

    RegistroCreate, RegistroOut,

    AccionMejoraCreate, AccionMejoraOut, PlanAccionUpdate, VerificacionUpdate, AccionMejoraEdicion,

    SolicitudCreate, SolicitudOut, SolicitudResponder, SolicitudEdicion,

    LoginRequest, GoogleLoginRequest, LoginResponse, UsuarioOut, UsuarioCreate, UsuarioEdicion,

    RecomendacionOut, LoteAccionesCreate, ObservacionCreate, AreaCreate, EvidenciaAsignar, NotificacionesOut,

)



Base.metadata.create_all(bind=engine)



def _ciclo_revision_periodica():

    """Cada hora revisa plazos (pasa a 'Pendiente de evidencia' y manda recordatorios) sin que nadie tenga que abrir el sistema."""

    while True:

        time.sleep(3600)

        db = SessionLocal()

        try:

            _revisar_vencidas_y_notificar(db)

            _revisar_solicitudes_vencidas(db)

        except Exception as e:

            print(f"[REVISIÓN PERIÓDICA] Error: {e}")

        finally:

            db.close()





@asynccontextmanager

async def lifespan(app: FastAPI):

    threading.Thread(target=_ciclo_revision_periodica, daemon=True).start()

    yield





app = FastAPI(title="Sistema de Registro y Seguimiento - API", lifespan=lifespan)



app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_origin_regex=r"https://[A-Za-z0-9.-]+\.(devtunnels\.ms|app\.github\.dev)",
    allow_methods=["*"],
    allow_headers=["*"],
)





# =========================================================================

# Autenticación

# =========================================================================



@app.post("/api/auth/login", response_model=LoginResponse)

def login(data: LoginRequest, db: Session = Depends(get_db)):

    usuario = db.query(Usuario).filter(Usuario.correo == data.correo).first()

    if not usuario or not verificar_password(data.password, usuario.password_hash):

        raise HTTPException(status_code=401, detail="Correo o contraseña incorrectos")

    if not usuario.activo:

        raise HTTPException(status_code=401, detail="Este usuario está desactivado")

    token = crear_token(usuario.correo, usuario.rol, usuario.sede)

    return LoginResponse(access_token=token, usuario=usuario)





# ID de cliente OAuth de Google (Google Cloud Console → Credenciales → ID de cliente OAuth 2.0,

# tipo "Aplicación web"). Ponerlo como variable de entorno GOOGLE_CLIENT_ID en el servidor.

GOOGLE_CLIENT_ID = os.environ.get("GOOGLE_CLIENT_ID","349423893149-r827m4h6cof4l312vd0n356l9ob0n73v.apps.googleusercontent.com")

DOMINIO_PERMITIDO = "unifranz.edu.bo"





@app.post("/api/auth/google", response_model=LoginResponse)

def login_google(data: GoogleLoginRequest, db: Session = Depends(get_db)):

    if not GOOGLE_CLIENT_ID:

        raise HTTPException(status_code=500, detail="El login con Google no está configurado en el servidor (falta GOOGLE_CLIENT_ID).")

    try:

        payload = google_id_token.verify_oauth2_token(

            data.credential, google_requests.Request(), GOOGLE_CLIENT_ID

        )

    except ValueError:

        raise HTTPException(status_code=401, detail="Token de Google inválido o vencido")



    correo = (payload.get("email") or "").lower().strip()

    if not payload.get("email_verified") or not correo.endswith(f"@{DOMINIO_PERMITIDO}"):

        raise HTTPException(status_code=403, detail=f"Solo se permite el ingreso con correos @{DOMINIO_PERMITIDO}")



    usuario = db.query(Usuario).filter(Usuario.correo == correo).first()

    if not usuario:

        # Primer ingreso: se crea la cuenta sola, como Director/a de Carrera sin área todavía.

        # Sistemas de Gestión le asigna el área desde Usuarios para que empiece a ver sus acciones.

        usuario = Usuario(

            nombre_completo=payload.get("name") or correo.split("@")[0],

            correo=correo,

            password_hash=hashear_password(uuid.uuid4().hex),  # nunca se usa: solo entra por Google

            rol="usuario",

            sede="El Alto",

            area=None,

            activo=True,

        )

        db.add(usuario)

        db.commit()

        db.refresh(usuario)

        _notificar_admins(

            db,

            "Nueva cuenta creada por Google — falta asignarle área",

            f"{usuario.nombre_completo} ({usuario.correo}) ingresó por primera vez con su correo institucional. "

            "Asígnale su Área/Carrera en Usuarios para que pueda ver sus Acciones de Mejora.",

        )

    elif not usuario.activo:

        raise HTTPException(status_code=401, detail="Este usuario está desactivado")



    token = crear_token(usuario.correo, usuario.rol, usuario.sede)

    return LoginResponse(access_token=token, usuario=usuario)





@app.get("/api/auth/me", response_model=UsuarioOut)

def me(usuario: Usuario = Depends(obtener_usuario_actual)):

    return usuario





@app.get("/api/usuarios", response_model=List[UsuarioOut])

def listar_usuarios(db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    return db.query(Usuario).order_by(Usuario.nombre_completo).all()





@app.post("/api/usuarios", response_model=UsuarioOut)

def crear_usuario(data: UsuarioCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    correo_normalizado = (data.correo or "").strip().lower()

    if not correo_normalizado.endswith(f"@{DOMINIO_PERMITIDO}"):

        raise HTTPException(status_code=400, detail=f"El correo debe ser institucional (@{DOMINIO_PERMITIDO}).")

    data.correo = correo_normalizado

    existente = db.query(Usuario).filter(Usuario.correo == data.correo).first()

    if existente:

        raise HTTPException(status_code=400, detail="Ya existe un usuario con ese correo.")

    nuevo = Usuario(

        nombre_completo=data.nombre_completo,

        correo=data.correo,

        password_hash=hashear_password(data.password),

        rol=data.rol,

        sede=data.sede if data.rol == "usuario" else None,

        activo=True,

    )

    db.add(nuevo)

    db.flush()

    if data.rol == "usuario":

        asignar_areas(db, nuevo, data.areas or parse_areas(data.area))

    db.commit()

    db.refresh(nuevo)

    return nuevo





@app.patch("/api/usuarios/{usuario_id}", response_model=UsuarioOut)

def editar_usuario(usuario_id: int, data: UsuarioEdicion, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    objetivo = db.query(Usuario).filter(Usuario.id == usuario_id).first()

    if not objetivo:

        raise HTTPException(status_code=404, detail="Usuario no encontrado")



    cambios = data.dict(exclude_unset=True)

    password_nueva = cambios.pop("password", None)

    areas_nuevas = cambios.pop("areas", None)

    area_texto = cambios.pop("area", None)

    for key, value in cambios.items():

        setattr(objetivo, key, value)

    if password_nueva:

        objetivo.password_hash = hashear_password(password_nueva)

    if objetivo.rol == "usuario":

        if areas_nuevas is not None:

            asignar_areas(db, objetivo, areas_nuevas)

        elif area_texto is not None:

            asignar_areas(db, objetivo, parse_areas(area_texto))



    db.commit()

    db.refresh(objetivo)

    return objetivo





# --- Notificaciones: buscan a quién avisar según sede/rol, mandan el correo y dejan el aviso en la campanita ---

PATRON_NUMERO_ACCION = re.compile(r"\bAM[ -][A-Z0-9]+(?:-[A-Z0-9]+)*")


def _accion_mencionada(db: Session, *textos: str) -> Optional[int]:
    """Devuelve el id de una única Acción de Mejora mencionada en el asunto o mensaje."""
    for texto in textos:
        numeros = set(PATRON_NUMERO_ACCION.findall(texto or ""))
        if not numeros:
            continue
        ids = [a.id for a in db.query(AccionMejora).filter(AccionMejora.numero.in_(numeros)).all()]
        if len(ids) == 1:
            return ids[0]
    return None


def _avisar(db: Session, usuario: Usuario, asunto: str, mensaje: str):
    """Un aviso a una persona: correo + campanita dentro del sistema."""
    enviar_correo(usuario.correo, asunto, mensaje)
    db.add(Notificacion(
        usuario_id=usuario.id,
        asunto=asunto,
        mensaje=mensaje,
        fecha=_ahora(),
        leida=False,
        accion_id=_accion_mencionada(db, asunto, mensaje),
    ))
    db.commit()


def _notificar_usuarios_de_sede(db: Session, sede: Optional[str], asunto: str, mensaje: str):
    """Avisa a todos los usuarios (Directores de Carrera) activos de esa sede."""
    if not sede:
        return
    destinatarios = db.query(Usuario).filter(Usuario.sede == sede, Usuario.rol == "usuario", Usuario.activo == True).all()
    for u in destinatarios:
        _avisar(db, u, asunto, mensaje)


def _norm_area(valor: Optional[str]) -> str:
    return (valor or "").strip().lower()


def _notificar_usuarios_de_area(db: Session, area, asunto: str, mensaje: str):
    """Avisa una sola vez a cada Director activo que pertenezca a alguna de las áreas."""
    nombres = [area] if isinstance(area, str) else list(area or [])
    objetivo = {_norm_area(n) for n in nombres if _norm_area(n)}
    if not objetivo:
        return
    candidatos = db.query(Usuario).filter(Usuario.rol == "usuario", Usuario.activo == True).all()
    for u in candidatos:
        if objetivo & {_norm_area(a) for a in u.areas}:
            _avisar(db, u, asunto, mensaje)


def _notificar_admins(db: Session, asunto: str, mensaje: str):
    """Avisa a todos los administradores (Sistemas de Gestión) activos."""
    admins = db.query(Usuario).filter(Usuario.rol == "admin", Usuario.activo == True).all()
    for a in admins:
        _avisar(db, a, asunto, mensaje)


# --- Campanita: cada persona ve y marca solo SUS avisos ---

@app.get("/api/notificaciones", response_model=NotificacionesOut)
def mis_notificaciones(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):
    propias = db.query(Notificacion).filter(Notificacion.usuario_id == usuario.id)
    return {
        "no_leidas": propias.filter(Notificacion.leida == False).count(),
        "items": propias.order_by(Notificacion.id.desc()).limit(50).all(),
    }


@app.post("/api/notificaciones/leer-todas")
def marcar_todas_leidas(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):
    db.query(Notificacion).filter(
        Notificacion.usuario_id == usuario.id, Notificacion.leida == False
    ).update({"leida": True})
    db.commit()
    return {"ok": True}


@app.post("/api/notificaciones/{notificacion_id}/leer")
def marcar_leida(notificacion_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):
    n = db.query(Notificacion).filter(
        Notificacion.id == notificacion_id, Notificacion.usuario_id == usuario.id
    ).first()
    if not n:
        raise HTTPException(status_code=404, detail="Notificación no encontrada")
    n.leida = True
    db.commit()
    return {"ok": True}


# =========================================================================

# Registros — módulo general de incidencias / oportunidades de mejora.

# NOTA: esta tabla todavía no tiene campo "sede", así que no se puede

# filtrar por sede todavía (solo exige estar logueado).

# =========================================================================



@app.get("/api/registros", response_model=List[RegistroOut])

def listar_registros(q: Optional[str] = None, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    query = db.query(Registro)

    if q:

        like = f"%{q}%"

        query = query.filter(

            or_(

                Registro.titulo.ilike(like),

                Registro.sistema.ilike(like),

                Registro.categoria.ilike(like),

                Registro.tipo.ilike(like),

            )

        )

    return query.order_by(Registro.id.desc()).all()





@app.get("/api/registros/{registro_id}/relacionados", response_model=List[RegistroOut])

def registros_relacionados(registro_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    base = db.query(Registro).filter(Registro.id == registro_id).first()

    if not base:

        raise HTTPException(status_code=404, detail="Registro no encontrado")

    return (

        db.query(Registro)

        .filter(

            Registro.id != registro_id,

            or_(Registro.sistema == base.sistema, Registro.categoria == base.categoria),

        )

        .all()

    )





@app.post("/api/registros", response_model=RegistroOut)

def crear_registro(registro: RegistroCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    nuevo = Registro(**registro.dict())

    db.add(nuevo)

    db.commit()

    db.refresh(nuevo)

    return nuevo





@app.get("/api/dashboard")

def resumen_dashboard(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    registros = db.query(Registro).all()

    por_sistema, por_estado = {}, {}

    for r in registros:

        por_sistema[r.sistema] = por_sistema.get(r.sistema, 0) + 1

        por_estado[r.estado] = por_estado.get(r.estado, 0) + 1

    return {

        "total": len(registros),

        "incidencias": sum(1 for r in registros if r.tipo == "Incidencia"),

        "mejoras": sum(1 for r in registros if r.tipo == "Oportunidad de mejora"),

        "pendientes": sum(1 for r in registros if r.estado != "Resuelto"),

        "por_sistema": [{"sistema": k, "cantidad": v} for k, v in por_sistema.items()],

        "por_estado": [{"estado": k, "cantidad": v} for k, v in por_estado.items()],

    }





# =========================================================================

# Acciones de Mejora — flujo de 3 etapas (formulario CAPA de calidad)

# =========================================================================



def _agregar_historial(db: Session, accion: AccionMejora, evento: str):

    db.add(HistorialAccion(accion_id=accion.id, fecha=_ahora(), evento=evento))





def _ahora() -> str:

    return datetime.now().strftime("%Y-%m-%d %H:%M")





def _verificar_acceso_accion(usuario: Usuario, accion: AccionMejora):

    """El Director/a solo accede a las acciones de SU área/carrera. El admin ve todo."""

    if usuario.rol == "admin":

        return

    mias = {a.id for a in usuario.areas_rel}

    if not mias & {a.id for a in accion.areas_rel}:

        raise HTTPException(status_code=403, detail="Esta acción no corresponde a tu área.")





def _verificar_acceso_sede(usuario: Usuario, sede: Optional[str]):

    """Un usuario (Director/a de Carrera) solo puede acceder a lo de su propia sede. El admin ve todo."""

    if usuario.rol != "admin" and (sede or "") != (usuario.sede or ""):

        raise HTTPException(status_code=403, detail="No tenés acceso a un registro de otra sede.")





@app.get("/api/acciones-mejora", response_model=List[AccionMejoraOut])

def listar_acciones(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    _revisar_vencidas_y_notificar(db)

    query = db.query(AccionMejora).options(joinedload(AccionMejora.historial))

    if usuario.rol != "admin":

        ids = [a.id for a in usuario.areas_rel]

        if not ids:

            return []

        query = query.filter(AccionMejora.areas_rel.any(Area.id.in_(ids)))

    return query.order_by(AccionMejora.id.desc()).all()





FILTROS_EXPORTACION = {

    "todos": "Todos los registros",

    "anual": "Solo un año específico",

    "pendientes_entrega": "Pendientes de entrega del plan de acción",

    "en_ejecucion": "En ejecución del plan",

    "pendientes_evidencia": "Pendientes de evidencia",

    "pendientes_verificacion": "Pendientes de verificación",

    "cerrados": "Cerrados",

}



COLUMNAS_DISPONIBLES = {

    "numero": ("Número", lambda r: r.numero),

    "numero_ac_original": ("N° AC Original", lambda r: r.numero_ac_original),

    "sede": ("Sede", lambda r: r.sede),

    "estado": ("Estado", lambda r: r.estado),

    "estado_original": ("Estado Original (Excel)", lambda r: r.estado_original),

    "solicitante": ("Solicitante", lambda r: r.solicitante),

    "area_responsable": ("Área Responsable", lambda r: r.area_responsable),

    "categoria": ("Categoría", lambda r: r.categoria),

    "subcategoria": ("Subcategoría", lambda r: r.subcategoria),

    "declaracion_hallazgo": ("Hallazgo", lambda r: r.declaracion_hallazgo),

    "fuente_identificacion": ("Fuente de Identificación", lambda r: r.fuente_identificacion),

    "procesos": ("Proceso(s)", lambda r: r.procesos),

    "requiere_cambio_sgc": ("¿Requiere cambio SGC?", lambda r: r.requiere_cambio_sgc),

    "fecha_creacion": ("Fecha de Creación", lambda r: r.fecha_creacion),

    "plazo_ejecucion": ("Plazo de Ejecución", lambda r: r.plazo_ejecucion),

    "fecha_cierre_plan": ("Fecha de Cierre (Excel)", lambda r: r.fecha_cierre_plan),

    "fecha_verificacion": ("Fecha de Verificación", lambda r: r.fecha_verificacion),

    "plan_accion": ("Plan de Acción", lambda r: r.plan_accion),

    "conclusion": ("Conclusión", lambda r: r.conclusion),

    "observacion": ("Observación", lambda r: r.observacion),

}





@app.get("/api/acciones-mejora/exportar")

def exportar_acciones(

    filtro: str = "todos",

    anio: Optional[str] = None,

    ids: Optional[str] = None,

    columnas: Optional[str] = None,

    db: Session = Depends(get_db),

    usuario: Usuario = Depends(requerir_admin_descarga),

):

    query = db.query(AccionMejora)



    if ids:

        id_list = [int(x) for x in ids.split(",") if x.strip().isdigit()]

        query = query.filter(AccionMejora.id.in_(id_list))

    elif filtro == "pendientes_entrega":

        query = query.filter(AccionMejora.estado == "Pendiente Plan de Acción")

    elif filtro == "en_ejecucion":

        query = query.filter(AccionMejora.estado == "En ejecución")

    elif filtro == "pendientes_evidencia":

        query = query.filter(AccionMejora.estado == "Pendiente de evidencia")

    elif filtro == "pendientes_verificacion":

        query = query.filter(AccionMejora.estado == "Pendiente Verificación")

    elif filtro == "cerrados":

        query = query.filter(AccionMejora.estado == "Cerrada")

    elif filtro == "anual" and anio:

        query = query.filter(AccionMejora.fecha_creacion.like(f"{anio}%"))



    registros = query.order_by(AccionMejora.id.desc()).all()



    columnas_elegidas = [c for c in (columnas.split(",") if columnas else []) if c in COLUMNAS_DISPONIBLES]

    if not columnas_elegidas:

        columnas_elegidas = list(COLUMNAS_DISPONIBLES.keys())



    wb = Workbook()

    ws = wb.active

    ws.title = "Acciones de Mejora"



    encabezados = [COLUMNAS_DISPONIBLES[c][0] for c in columnas_elegidas]

    ws.append(encabezados)

    for cell in ws[1]:

        cell.font = Font(bold=True, color="FFFFFF")

        cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")



    for r in registros:

        ws.append([COLUMNAS_DISPONIBLES[c][1](r) for c in columnas_elegidas])



    for col in ws.columns:

        largo = max((len(str(c.value)) for c in col if c.value), default=10)

        ws.column_dimensions[col[0].column_letter].width = min(largo + 2, 45)



    buffer = BytesIO()

    wb.save(buffer)

    buffer.seek(0)



    nombre_archivo = f"acciones_mejora_{'seleccion' if ids else filtro}.xlsx"

    return StreamingResponse(

        buffer,

        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},

    )





MAPA_ESTADOS_EXCEL = {

    "Llenado de formulario pendiente": "Pendiente Plan de Acción",

    "Dentro de plazo": "Pendiente Plan de Acción",

    "Debe cerrarse el presente mes": "Pendiente Plan de Acción",

    "Fuera de Plazo": "Pendiente Plan de Acción",

    "Fuera de plazo": "Pendiente Plan de Acción",

    "Pendiente la verificación de la eficacia": "Pendiente Verificación",

    "Cerrada": "Cerrada",

    "Apertura de una nueva acción (Plan de Acción No Implementado)": "Pendiente Plan de Acción",

    "Apertura de una nueva acción (Plan de Accion Implementado No Eficaz)": "Pendiente Plan de Acción",

}



COLUMNAS_ESPERADAS = {"Origen", "N° de AC"}



MAPA_COLUMNAS = {

    "Origen": "origen",

    "N° de AC": "numero_ac",

    "Descripción": "descripcion",

    "Plan de Acción": "plan",

    "Categoría": "categoria",

    "Subcategoría": "subcategoria",

    "Área Responsable": "area",

    "¿Se requiere un cambio en el SGC?": "requiere_cambio",

    "Plazo del Plan de Acción": "plazo",

    "Fecha de Cierre Plan de Acción": "fecha_cierre",

    "Estado": "estado_txt",

    "Observación": "observacion",

}





def _formatear_fecha(valor):

    if valor is None:

        return None

    if hasattr(valor, "isoformat"):

        return valor.isoformat()[:10]

    return str(valor).strip() or None





@app.post("/api/acciones-mejora/importar-excel")

async def importar_excel(archivo: UploadFile = File(...), db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    contenido = await archivo.read()

    try:

        wb = load_workbook(BytesIO(contenido), data_only=True)

    except Exception:

        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. ¿Es un .xlsx válido?")



    hojas_procesadas = []

    importados = 0



    for nombre_hoja in wb.sheetnames:

        ws = wb[nombre_hoja]



        fila_encabezado = None

        indices = {}

        for r in range(1, min(6, ws.max_row + 1)):

            valores = {c.value: c.column for c in ws[r] if c.value}

            if COLUMNAS_ESPERADAS.issubset(valores.keys()):

                fila_encabezado = r

                indices = {

                    campo: valores[encabezado]

                    for encabezado, campo in MAPA_COLUMNAS.items()

                    if encabezado in valores

                }

                break

        if fila_encabezado is None:

            continue



        sede = nombre_hoja.replace("Sede", "").strip() or "Nacional"

        filas_de_esta_hoja = 0



        def valor(fila, campo):

            idx = indices.get(campo)

            return fila[idx - 1] if idx and idx <= len(fila) else None



        for fila in ws.iter_rows(min_row=fila_encabezado + 1, max_row=ws.max_row, values_only=True):

            descripcion = valor(fila, "descripcion")

            if not descripcion:

                continue



            estado_txt = valor(fila, "estado_txt")

            total = db.query(AccionMejora).count()

            numero_interno = f"AM-{total + 1:04d}"

            numero_ac = valor(fila, "numero_ac")



            nueva = AccionMejora(

                numero=numero_interno,

                numero_ac_original=str(numero_ac) if numero_ac else None,

                sede=sede,

                estado=MAPA_ESTADOS_EXCEL.get(estado_txt, "Pendiente Plan de Acción"),

                estado_original=estado_txt,

                declaracion_hallazgo=str(descripcion),

                fuente_identificacion=valor(fila, "origen"),

                categoria=valor(fila, "categoria"),

                subcategoria=valor(fila, "subcategoria"),

                plan_accion=valor(fila, "plan"),

                requiere_cambio_sgc=valor(fila, "requiere_cambio"),

                plazo_ejecucion=_formatear_fecha(valor(fila, "plazo")),

                fecha_cierre_plan=_formatear_fecha(valor(fila, "fecha_cierre")),

                observacion=valor(fila, "observacion"),

                fecha_creacion=_ahora(),

            )

            db.add(nueva)

            db.flush()

            area_txt = valor(fila, "area")

            asignar_areas(db, nueva, parse_areas(str(area_txt)) if area_txt else [])

            db.commit()

            db.refresh(nueva)

            _agregar_historial(

                db, nueva,

                f"Importado desde Excel (hoja: {nombre_hoja}, estado original: {estado_txt or 'sin estado'})",

            )

            db.commit()

            importados += 1

            filas_de_esta_hoja += 1



        if filas_de_esta_hoja > 0:

            hojas_procesadas.append(f"{nombre_hoja} ({filas_de_esta_hoja} registros)")



    if not hojas_procesadas:

        raise HTTPException(

            status_code=400,

            detail="No se encontró ninguna hoja con el formato esperado (columnas 'Origen' y 'N° de AC').",

        )



    return {"importados": importados, "hojas_procesadas": hojas_procesadas}





@app.get("/api/acciones-mejora/{accion_id}", response_model=AccionMejoraOut)

def obtener_accion(accion_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    accion = (

        db.query(AccionMejora)

        .options(joinedload(AccionMejora.historial))

        .filter(AccionMejora.id == accion_id)

        .first()

    )

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    _verificar_acceso_accion(usuario, accion)

    return accion





@app.patch("/api/acciones-mejora/{accion_id}", response_model=AccionMejoraOut)

def editar_accion(accion_id: int, data: AccionMejoraEdicion, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")



    cambios = data.dict(exclude_unset=True)

    areas_nuevas = cambios.pop("areas", None)

    if areas_nuevas is not None:

        cambios.pop("area_responsable", None)

        asignar_areas(db, accion, _nombres_areas(areas_nuevas))

    elif "area_responsable" in cambios:

        asignar_areas(db, accion, _nombres_areas(None, cambios.pop("area_responsable")))

    for key, value in cambios.items():

        setattr(accion, key, value)



    if cambios or areas_nuevas is not None:

        db.commit()

        _agregar_historial(db, accion, f"Editado manualmente ({', '.join(list(cambios.keys()) + (['areas'] if areas_nuevas is not None else []))})")

        db.commit()

    db.refresh(accion)

    return accion





SEDE_CODIGOS = {

    "El Alto": "EA",

    "Cochabamba": "CB",

    "Santa Cruz": "SC",

    "La Paz": "LP",

    "Nacional": "NA",

}

SEDE_DEFECTO = "El Alto"

PLAZOS_ENTREGA_DIAS = (7, 14)





def _codigo_sede(sede: Optional[str]) -> str:

    return SEDE_CODIGOS.get((sede or "").strip() or SEDE_DEFECTO, "EA")





def _generar_numero_ac(db: Session, origen: Optional[str], sede: Optional[str] = None) -> str:

    """

    Código correlativo de la Acción de Mejora: AM {FUENTE}-{SEDE}-{N}-{AA}, ej. AM BS-EA-14-26

    (BS = Buzón de Sugerencias, EA = El Alto). N es el número MÁS ALTO existente + 1 (no una

    cuenta), y considera tanto los códigos del sistema como los originales importados de Excel,

    con o sin código de sede (el formato viejo era AM BS-14-26).

    """

    origen_code = (origen or "BS").strip().upper() or "BS"

    sede_code = _codigo_sede(sede)

    anio = datetime.now().strftime("%y")

    patron = re.compile(rf"^AM\s*{re.escape(origen_code)}-(?:{sede_code}-)?(\d+)-{anio}$", re.IGNORECASE)

    maximo = 0

    for numero, original in db.query(AccionMejora.numero, AccionMejora.numero_ac_original).all():

        for valor in (numero, original):

            if valor:

                m = patron.match(valor.strip())

                if m:

                    maximo = max(maximo, int(m.group(1)))

    return f"AM {origen_code}-{sede_code}-{maximo + 1:02d}-{anio}"





def _plazo_desde_dias(dias: Optional[int]) -> Optional[str]:

    """7 o 14 días desde hoy, como fecha ISO. None si no se indicó."""

    if dias is None:

        return None

    if dias not in PLAZOS_ENTREGA_DIAS:

        raise HTTPException(status_code=400, detail="El plazo para entregar el plan debe ser de 7 o 14 días.")

    return (date.today() + timedelta(days=dias)).isoformat()





def _nombres_areas(areas: Optional[List[str]], texto: Optional[str] = None) -> List[str]:

    """Lista de áreas limpia: usa `areas` si viene, si no separa el texto ('ODO / DAF')."""

    base = areas if areas else parse_areas(texto)

    return parse_areas(";".join(a for a in base if a and a.strip()))





@app.post("/api/acciones-mejora", response_model=AccionMejoraOut)

def crear_accion(data: AccionMejoraCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    """Registro de una Acción de Mejora. Solo admin. Si se completan directamente

    el Plan de Acción o el Estado (como al transcribir un caso ya resuelto del

    Excel), se respetan tal cual; si no, arranca en la etapa inicial normal."""

    datos = data.dict()

    nombres_areas = _nombres_areas(datos.pop("areas", None), datos.get("area_responsable"))

    plazo_dias = datos.pop("plazo_dias", None)

    if plazo_dias is not None:

        datos["plazo_entrega_plan"] = _plazo_desde_dias(plazo_dias)

    datos["sede"] = (datos.get("sede") or "").strip() or SEDE_DEFECTO

    datos.pop("area_responsable", None)



    numero = _generar_numero_ac(db, data.fuente_identificacion, datos["sede"])

    estado_inicial = datos.pop("estado", None) or "Pendiente Plan de Acción"

    accion = AccionMejora(

        numero=numero,

        estado=estado_inicial,

        fecha_creacion=_ahora(),

        **datos,

    )

    db.add(accion)

    db.flush()

    asignar_areas(db, accion, nombres_areas)

    db.commit()

    db.refresh(accion)

    _agregar_historial(db, accion, "Registrada")

    db.commit()

    if estado_inicial == "Pendiente Plan de Acción":

        _notificar_usuarios_de_area(

            db, accion.areas,

            f"Nueva Acción de Mejora asignada — {numero}",

            f"Se registró la Acción de Mejora {numero} y está pendiente de tu plan de acción.\n\nHallazgo: {accion.declaracion_hallazgo}"

            + (f"\n\nPlazo para entregar el plan de acción: {accion.plazo_entrega_plan}" if accion.plazo_entrega_plan else ""),

        )

    db.refresh(accion)

    return accion





@app.post("/api/acciones-mejora/lote")

def crear_lote_acciones(data: LoteAccionesCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    """

    Una solicitud del buzón de sugerencias para UNA o VARIAS áreas: crea una Acción de Mejora por

    cada hallazgo (cada una con su código AM BS-EA-NN-AA correlativo, compartida por esas áreas) y

    guarda las recomendaciones. Cada persona de las áreas recibe UN SOLO correo con todo.

    """

    areas = _nombres_areas(data.areas, data.area)

    if not areas:

        raise HTTPException(status_code=400, detail="Indica el área o carrera destinataria.")

    hallazgos = [h.strip() for h in data.hallazgos if h and h.strip()]

    recomendaciones = [r.strip() for r in data.recomendaciones if r and r.strip()]

    if not hallazgos and not recomendaciones:

        raise HTTPException(status_code=400, detail="Agrega al menos una acción de mejora o una recomendación.")



    plazo = _plazo_desde_dias(data.plazo_dias) if data.plazo_dias is not None else (data.plazo_entrega_plan or None)



    usuarios_activos = db.query(Usuario).filter(Usuario.rol == "usuario", Usuario.activo == True).all()

    for nombre in areas:

        if not any(_norm_area(nombre) in {_norm_area(a) for a in u.areas} for u in usuarios_activos):

            raise HTTPException(

                status_code=400,

                detail=f"No hay ningún usuario activo asignado al área '{nombre}'. Créalo primero en Usuarios.",

            )



    creadas = []

    for texto in hallazgos:

        accion = AccionMejora(

            numero=_generar_numero_ac(db, "BS", SEDE_DEFECTO),

            estado="Pendiente Plan de Acción",

            fecha_creacion=_ahora(),

            declaracion_hallazgo=texto,

            fuente_identificacion="BS",

            solicitante=(data.solicitante or "").strip() or None,

            plazo_entrega_plan=plazo,

            sede=SEDE_DEFECTO,

        )

        db.add(accion)

        db.flush()  # para que el siguiente código correlativo ya cuente este

        asignar_areas(db, accion, areas)

        _agregar_historial(db, accion, "Registrada y enviada al área")

        creadas.append(accion)

    nombre_areas = " / ".join(areas)

    for texto in recomendaciones:

        for nombre in areas:

            db.add(Recomendacion(area=nombre, texto=texto, fecha=_ahora(), enviada_por=usuario.nombre_completo))

    db.commit()



    partes = [f"Sistemas de Gestión te envió una solicitud del Buzón de Sugerencias para: {nombre_areas}."]

    if creadas:

        partes.append("\nAcciones de Mejora (pendientes de tu plan de acción):")

        for a in creadas:

            partes.append(f"  • {a.numero}: {a.declaracion_hallazgo}")

        if plazo:

            partes.append(f"\nPlazo para entregar el plan de acción: {plazo}")

    if recomendaciones:

        partes.append("\nRecomendaciones (para tu conocimiento):")

        for r in recomendaciones:

            partes.append(f"  • {r}")

    partes.append("\nIngresa al sistema para completar tus acciones.")

    _notificar_usuarios_de_area(

        db, areas,

        f"Nueva solicitud del Buzón de Sugerencias — {len(creadas)} acción(es), {len(recomendaciones)} recomendación(es)",

        "\n".join(partes),

    )

    return {

        "acciones": [{"id": a.id, "numero": a.numero} for a in creadas],

        "recomendaciones": len(recomendaciones),

        "plazo_entrega_plan": plazo,

    }





@app.get("/api/recomendaciones", response_model=List[RecomendacionOut])

def listar_recomendaciones(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    query = db.query(Recomendacion)

    if usuario.rol != "admin":

        mias = [_norm_area(a) for a in usuario.areas]

        if not mias:

            return []

        query = query.filter(func.lower(func.trim(Recomendacion.area)).in_(mias))

    return query.order_by(Recomendacion.id.desc()).all()





@app.get("/api/areas", response_model=List[str])

def listar_areas(db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    """Catálogo de áreas/carreras activas (para asignarlas a usuarios y a solicitudes)."""

    return [a.nombre for a in db.query(Area).filter(Area.activa == True).order_by(Area.nombre).all()]





@app.post("/api/areas", response_model=List[str])

def crear_area(data: AreaCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    nombre = (data.nombre or "").strip()

    if not nombre:

        raise HTTPException(status_code=400, detail="Escribe el nombre del área.")

    area = obtener_o_crear_area(db, nombre)

    area.activa = True

    db.commit()

    return [a.nombre for a in db.query(Area).filter(Area.activa == True).order_by(Area.nombre).all()]





@app.post("/api/acciones-mejora/{accion_id}/observacion", response_model=AccionMejoraOut)

def observar_accion(accion_id: int, data: ObservacionCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    """La ingeniera deja un comentario para que el área corrija: estado 'Observada' + correo al área."""

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    texto = (data.texto or "").strip()

    if not texto:

        raise HTTPException(status_code=400, detail="Escribe el comentario para el área.")

    if accion.estado not in ("En ejecución", "Pendiente de evidencia", "Pendiente Verificación"):

        raise HTTPException(

            status_code=400,

            detail="Solo se puede observar una acción cuyo plan de acción ya fue entregado y que todavía no está cerrada.",

        )



    accion.estado = "Observada"

    accion.ultima_observacion = texto

    db.commit()

    _agregar_historial(db, accion, f"Observación de Sistemas de Gestión: {texto}")

    db.commit()

    _notificar_usuarios_de_area(

        db, accion.areas,

        f"Observación en tu Acción de Mejora — {accion.numero}",

        f"Sistemas de Gestión dejó una observación en la acción {accion.numero}:\n\n{texto}\n\n"

        "Ingresa al sistema, corrige lo indicado y vuelve a enviar.",

    )

    db.refresh(accion)

    return accion





# --- Evidencias (archivos PDF / Excel / fotos) ---



UPLOAD_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads", "evidencias")

EXTENSIONES_PERMITIDAS = {".pdf", ".xls", ".xlsx", ".csv", ".doc", ".docx", ".jpg", ".jpeg", ".png", ".webp"}

TAMANO_MAX_BYTES = 15 * 1024 * 1024





def _checklist_evidencias(accion: AccionMejora):

    """Cruza lo que el área prometió en el formulario con los archivos subidos.

    Devuelve (items_sin_evidencia, archivos_sin_asignar). Sin plan en lista, no hay nada que cruzar."""

    if not accion.plan_items:

        return [], []

    ids = {i.id for i in accion.plan_items}

    con_archivo = {ev.plan_item_id for ev in accion.evidencias if ev.plan_item_id in ids}

    faltan = [i for i in accion.plan_items if i.id not in con_archivo]

    sin_asignar = [ev for ev in accion.evidencias if ev.plan_item_id not in ids]

    return faltan, sin_asignar





def _texto_faltantes(faltan) -> str:

    return "; ".join(f"{i.orden + 1}. {i.accion_texto} (evidencia: {i.evidencia or '—'})" for i in faltan)





@app.post("/api/acciones-mejora/{accion_id}/evidencias", response_model=AccionMejoraOut)

async def subir_evidencia(

    accion_id: int,

    archivo: UploadFile = File(...),

    plan_item_id: Optional[int] = Form(None),

    db: Session = Depends(get_db),

    usuario: Usuario = Depends(obtener_usuario_actual),

):

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    _verificar_acceso_accion(usuario, accion)

    if plan_item_id is not None and plan_item_id not in {i.id for i in accion.plan_items}:

        raise HTTPException(status_code=400, detail="Esa acción del plan no pertenece a esta Acción de Mejora.")

    if plan_item_id is None and accion.plan_items and usuario.rol != "admin":

        raise HTTPException(status_code=400, detail="Indica a cuál acción del plan corresponde el archivo.")

    if usuario.rol != "admin":

        if accion.estado == "En ejecución":

            raise HTTPException(

                status_code=400,

                detail=f"La evidencia se sube cuando se cumpla el plazo de ejecución ({accion.plazo_ejecucion}).",

            )

        if accion.estado != "Pendiente de evidencia":

            raise HTTPException(status_code=400, detail="En este momento no se puede subir evidencia en esta acción.")



    extension = os.path.splitext(archivo.filename or "")[1].lower()

    if extension not in EXTENSIONES_PERMITIDAS:

        raise HTTPException(status_code=400, detail="Tipo de archivo no permitido. Usa PDF, Excel, Word o imágenes (JPG/PNG).")

    contenido = await archivo.read()

    if len(contenido) > TAMANO_MAX_BYTES:

        raise HTTPException(status_code=400, detail="El archivo supera el máximo de 15 MB.")



    os.makedirs(UPLOAD_DIR, exist_ok=True)

    nombre_guardado = f"{uuid.uuid4().hex}{extension}"

    with open(os.path.join(UPLOAD_DIR, nombre_guardado), "wb") as f:

        f.write(contenido)



    db.add(EvidenciaArchivo(

        accion_id=accion.id,

        plan_item_id=plan_item_id,

        nombre_original=os.path.basename(archivo.filename or nombre_guardado),

        nombre_guardado=nombre_guardado,

        fecha=_ahora(),

        subido_por=usuario.nombre_completo,

    ))

    db.commit()

    _agregar_historial(db, accion, f"Evidencia subida: {archivo.filename}")

    db.commit()

    db.refresh(accion)

    return accion





@app.post("/api/acciones-mejora/{accion_id}/enviar-evidencia", response_model=AccionMejoraOut)

def enviar_evidencia(accion_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    """El área terminó de subir sus archivos: la acción pasa a verificación de Sistemas de Gestión."""

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    _verificar_acceso_accion(usuario, accion)

    if accion.estado != "Pendiente de evidencia":

        raise HTTPException(status_code=400, detail="Esta acción no está esperando evidencia.")

    if not accion.evidencias:

        raise HTTPException(status_code=400, detail="Sube al menos un archivo de evidencia antes de enviar.")

    faltan, sin_asignar = _checklist_evidencias(accion)

    if faltan:

        raise HTTPException(status_code=400, detail=f"Falta la evidencia de: {_texto_faltantes(faltan)}")

    if sin_asignar:

        raise HTTPException(status_code=400, detail="Hay archivos sin asignar a una acción del plan; asígnalos antes de enviar.")



    accion.estado = "Pendiente Verificación"

    db.commit()

    _agregar_historial(db, accion, f"Evidencia enviada por {usuario.nombre_completo}: lista para verificación")

    db.commit()

    if accion.plan_items:

        bloques = []

        for it in accion.plan_items:

            nombres = [ev.nombre_original for ev in accion.evidencias if ev.plan_item_id == it.id]

            bloques.append(

                f"{it.orden + 1}. {it.accion_texto}\n   Evidencia indicada: {it.evidencia or '—'}\n"

                + "\n".join(f"   ✔ {n}" for n in nombres)

            )

        archivos = "\n\n".join(bloques)

    else:

        archivos = "\n".join(f"  • {ev.nombre_original}" for ev in accion.evidencias)

    _notificar_admins(

        db,

        f"Evidencia recibida — {accion.numero}",

        f"El área {accion.area_responsable or ''} envió la evidencia de la Acción de Mejora {accion.numero} "

        f"({len(accion.evidencias)} archivo(s)):\n{archivos}\n\n"

        "Revisa que esté completa y con los mismos nombres que indicaron en el formulario; "

        "luego completa la Segunda Parte (verificación) para cerrar la acción.",

    )

    # Confirmación para el área: que sepan que su evidencia sí llegó

    _notificar_usuarios_de_area(

        db, accion.areas,

        f"Recibimos tu evidencia — {accion.numero}",

        f"Recibimos la evidencia de la Acción de Mejora {accion.numero} ({len(accion.evidencias)} archivo(s)). "

        "Sistemas de Gestión la está verificando; te avisaremos cuando la acción quede cerrada "

        "o si necesita alguna corrección.",

    )

    db.refresh(accion)

    return accion





@app.patch("/api/evidencias/{evidencia_id}", response_model=AccionMejoraOut)

def asignar_evidencia(evidencia_id: int, data: EvidenciaAsignar, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    """Cambia a cuál acción del plan corresponde un archivo."""

    ev = db.query(EvidenciaArchivo).filter(EvidenciaArchivo.id == evidencia_id).first()

    if not ev:

        raise HTTPException(status_code=404, detail="Evidencia no encontrada")

    accion = db.query(AccionMejora).filter(AccionMejora.id == ev.accion_id).first()

    _verificar_acceso_accion(usuario, accion)

    if usuario.rol != "admin" and accion.estado != "Pendiente de evidencia":

        raise HTTPException(status_code=400, detail="Ya no se pueden reasignar archivos de esta acción.")

    if data.plan_item_id is not None and data.plan_item_id not in {i.id for i in accion.plan_items}:

        raise HTTPException(status_code=400, detail="Esa acción del plan no pertenece a esta Acción de Mejora.")

    ev.plan_item_id = data.plan_item_id

    db.commit()

    _agregar_historial(db, accion, f"Evidencia reasignada: {ev.nombre_original}")

    db.commit()

    db.refresh(accion)

    return accion





@app.get("/api/evidencias/{evidencia_id}/descargar")

def descargar_evidencia(evidencia_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_para_descarga)):

    ev = db.query(EvidenciaArchivo).filter(EvidenciaArchivo.id == evidencia_id).first()

    if not ev:

        raise HTTPException(status_code=404, detail="Evidencia no encontrada")

    accion = db.query(AccionMejora).filter(AccionMejora.id == ev.accion_id).first()

    _verificar_acceso_accion(usuario, accion)

    ruta = os.path.join(UPLOAD_DIR, ev.nombre_guardado)

    if not os.path.exists(ruta):

        raise HTTPException(status_code=404, detail="El archivo ya no está en el servidor")

    return FileResponse(ruta, filename=ev.nombre_original)





@app.delete("/api/evidencias/{evidencia_id}", response_model=AccionMejoraOut)

def eliminar_evidencia(evidencia_id: int, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    ev = db.query(EvidenciaArchivo).filter(EvidenciaArchivo.id == evidencia_id).first()

    if not ev:

        raise HTTPException(status_code=404, detail="Evidencia no encontrada")

    accion = db.query(AccionMejora).filter(AccionMejora.id == ev.accion_id).first()

    _verificar_acceso_accion(usuario, accion)

    if usuario.rol != "admin" and accion.estado != "Pendiente de evidencia":

        raise HTTPException(status_code=400, detail="Ya no se pueden quitar archivos de esta acción.")

    ruta = os.path.join(UPLOAD_DIR, ev.nombre_guardado)

    if os.path.exists(ruta):

        os.remove(ruta)

    nombre = ev.nombre_original

    db.delete(ev)

    db.commit()

    _agregar_historial(db, accion, f"Evidencia eliminada: {nombre}")

    db.commit()

    db.refresh(accion)

    return accion





def _estado_tras_plan(accion: AccionMejora) -> str:

    """

    Estado en el que queda la acción cuando el área entrega (o corrige) su plan:

    - plazo de ejecución todavía no cumplido  -> "En ejecución" (no se piden evidencias aún)

    - plazo cumplido y sin evidencia          -> "Pendiente de evidencia"

    - plazo cumplido y con evidencia          -> "Pendiente Verificación"

    """

    try:

        fin = date.fromisoformat(accion.plazo_ejecucion or "")

    except ValueError:

        return "Pendiente de evidencia"

    if fin > date.today():

        return "En ejecución"

    return "Pendiente Verificación" if accion.evidencias else "Pendiente de evidencia"





@app.put("/api/acciones-mejora/{accion_id}/plan-accion", response_model=AccionMejoraOut)

def completar_plan_accion(accion_id: int, data: PlanAccionUpdate, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    """Primera Parte del formulario — Sección a ser completada por el Área Responsable."""

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    _verificar_acceso_accion(usuario, accion)

    if accion.estado not in ("Pendiente Plan de Acción", "Observada"):

        raise HTTPException(status_code=400, detail="Esta acción no está pendiente de plan de acción")



    cambios = data.dict(exclude_unset=True)

    acciones_in = cambios.pop("acciones", None)

    plan_texto = cambios.pop("plan_accion", None)

    evidencia_texto = cambios.pop("evidencia_parte1", None)

    if acciones_in is None and (plan_texto or "").strip():

        acciones_in = [{"accion": plan_texto, "evidencia": evidencia_texto}]

    items = [

        ((i.get("accion") or "").strip(), (i.get("evidencia") or "").strip())

        for i in (acciones_in or [])

        if (i.get("accion") or "").strip()

    ]



    # --- Validaciones: la Primera Parte debe venir completa ---

    if not items:

        raise HTTPException(status_code=400, detail="Agrega al menos una acción al plan de acción.")

    if any(not ev for _, ev in items):

        raise HTTPException(status_code=400, detail="Indica la evidencia que presentarás por cada acción propuesta.")

    if not (cambios.get("analisis_causa_raiz") or "").strip():

        raise HTTPException(status_code=400, detail="Completa el análisis de causa raíz.")

    if not (cambios.get("responsables") or "").strip():

        raise HTTPException(status_code=400, detail="Indica el nombre del o los responsables.")

    try:

        plazo = date.fromisoformat((cambios.get("plazo_ejecucion") or "").strip())

    except ValueError:

        raise HTTPException(status_code=400, detail="Indica el plazo para la ejecución del plan de acción.")

    if plazo < date.today():

        raise HTTPException(status_code=400, detail="El plazo de ejecución no puede ser una fecha pasada.")



    era_observada = accion.estado == "Observada"

    for key, value in cambios.items():

        setattr(accion, key, value)

    # Se rehace la lista del plan: los archivos ya subidos quedan "sin asignar" hasta ligarlos de nuevo

    for ev in accion.evidencias:

        ev.plan_item_id = None

    accion.plan_items = [

        PlanAccionItem(orden=n, accion_texto=texto, evidencia=ev) for n, (texto, ev) in enumerate(items)

    ]

    # Texto de apoyo (exportaciones, búsqueda): las acciones y evidencias numeradas

    accion.plan_accion = "\n".join(f"{n + 1}. {t}" for n, (t, _) in enumerate(items))

    accion.evidencia_parte1 = "\n".join(f"{n + 1}. {ev}" for n, (_, ev) in enumerate(items))

    accion.fecha_llenado_parte1 = _ahora()

    accion.ultima_observacion = None

    accion.ultima_notificacion_evidencia = None

    db.flush()

    accion.estado = _estado_tras_plan(accion)

    db.commit()

    _agregar_historial(

        db, accion,

        ("Plan de acción corregido por el Área Responsable (tras observación)" if era_observada

         else "Plan de acción completado por el Área Responsable")

        + f" — plazo de ejecución: {accion.plazo_ejecucion}",

    )

    db.commit()

    corregido = " (corregido tras tu observación)" if era_observada else ""

    mensaje = (

        f"Ya te llegó el formulario lleno de la Acción de Mejora {accion.numero}{corregido}.\n\n"

        f"Área: {accion.area_responsable or 'sin área'}\n"

        f"Responsable(s): {accion.responsables or '—'}\n"

        f"Plazo de ejecución del plan: {accion.plazo_ejecucion}\n"

    )

    if accion.estado == "En ejecución":

        mensaje += "\nLas evidencias se le solicitarán automáticamente al área cuando llegue ese plazo."

    mensaje += "\n\nIngresa al sistema para revisar el formulario."

    _notificar_admins(db, f"Formulario lleno recibido — {accion.numero}", mensaje)

    # Confirmación para el área: que sepan que su formulario sí llegó

    confirmacion = (

        f"Recibimos tu formulario de la Acción de Mejora {accion.numero}"

        + (" (con las correcciones)" if era_observada else "")

        + f". Plazo de ejecución de tu plan: {accion.plazo_ejecucion}."

    )

    if accion.estado == "En ejecución":

        confirmacion += " Te avisaremos por correo cuando se acerque y llegue el día de subir tus evidencias."

    _notificar_usuarios_de_area(db, accion.areas, f"Recibimos tu formulario — {accion.numero}", confirmacion)

    db.refresh(accion)

    return accion





def _crear_accion_derivada(db: Session, original: AccionMejora) -> AccionMejora:

    numero = _generar_numero_ac(db, original.fuente_identificacion, original.sede)

    derivada = AccionMejora(

        numero=numero,

        estado="Pendiente Plan de Acción",

        fecha_creacion=_ahora(),

        declaracion_hallazgo=f"[Reincidencia de {original.numero}] {original.declaracion_hallazgo}",

        fuente_identificacion=original.fuente_identificacion,

        procesos=original.procesos,

        solicitante=original.solicitante,

        sede=original.sede,

    )

    db.add(derivada)

    db.flush()

    asignar_areas(db, derivada, original.areas)

    db.commit()

    db.refresh(derivada)

    _agregar_historial(

        db, derivada,

        f"Creada automáticamente porque la verificación de {original.numero} resultó ineficaz",

    )

    db.commit()

    _notificar_usuarios_de_area(

        db, derivada.areas,

        f"Nueva Acción de Mejora (reincidencia) — {derivada.numero}",

        f"Se abrió la Acción de Mejora {derivada.numero} porque {original.numero} no fue eficaz. Completa el plan de acción.",

    )

    return derivada





@app.put("/api/acciones-mejora/{accion_id}/verificacion", response_model=AccionMejoraOut)

def completar_verificacion(accion_id: int, data: VerificacionUpdate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    accion = db.query(AccionMejora).filter(AccionMejora.id == accion_id).first()

    if not accion:

        raise HTTPException(status_code=404, detail="Acción de Mejora no encontrada")

    if accion.estado != "Pendiente Verificación":

        raise HTTPException(status_code=400, detail="Esta acción no está pendiente de verificación")

    faltan, sin_asignar = _checklist_evidencias(accion)

    if faltan:

        raise HTTPException(status_code=400, detail=f"No se puede cerrar: falta la evidencia de {_texto_faltantes(faltan)}")

    if sin_asignar:

        raise HTTPException(status_code=400, detail="No se puede cerrar: hay archivos sin asignar a una acción del plan.")



    for key, value in data.dict(exclude_unset=True).items():

        setattr(accion, key, value)

    # Segunda Parte: fecha de verificación (hoy si no se indica) y nombre real de quien verifica

    accion.fecha_verificacion = (data.fecha_verificacion or "").strip() or _ahora()

    accion.coordinador_sistemas_gestion = (data.coordinador_sistemas_gestion or "").strip() or usuario.nombre_completo

    accion.estado = "Cerrada"



    evento = "Verificada y cerrada por Sistemas de Gestión (calidad)"

    if accion.es_ineficaz:

        nueva = _crear_accion_derivada(db, accion)

        accion.numero_accion_derivada = nueva.numero

        evento += f" — marcada como INEFICAZ, se abrió {nueva.numero} para continuar el seguimiento"



    db.commit()

    _agregar_historial(db, accion, evento)

    db.commit()

    # Aviso de cierre para el área

    cierre = f"La Acción de Mejora {accion.numero} fue verificada y quedó CERRADA por Sistemas de Gestión."

    if accion.es_ineficaz and accion.numero_accion_derivada:

        cierre += (

            f"\n\nLa verificación consideró que la acción no fue eficaz, por lo que se abrió "

            f"{accion.numero_accion_derivada} para continuar el seguimiento; recibirás otro correo con el detalle."

        )

    else:

        cierre += "\n\n¡Gracias por tu colaboración!"

    _notificar_usuarios_de_area(db, accion.areas, f"Acción de Mejora cerrada — {accion.numero}", cierre)

    db.refresh(accion)

    return accion





# Cuántos días antes del plazo se avisa al área. El día del plazo del formulario también se avisa.

AVISO_DIAS_ANTES = (3, 1)





def _revisar_vencidas_y_notificar(db: Session):

    """

    Revisión automática de plazos (se ejecuta cada hora y al abrir la lista):

      1. Plan de acción no entregado a tiempo      -> recordatorio diario al área.

      2. "En ejecución" y se cumplió el plazo       -> pasa a "Pendiente de evidencia" y avisa al área.

      3. "Pendiente de evidencia"                   -> recordatorio cada 3 días.

    """

    hoy = date.today()

    hoy_iso = hoy.isoformat()



    # 1) Plazo para entregar el plan vencido

    sin_plan = (

        db.query(AccionMejora)

        .filter(

            AccionMejora.estado == "Pendiente Plan de Acción",

            AccionMejora.plazo_entrega_plan.isnot(None),

            AccionMejora.plazo_entrega_plan != "",

        )

        .all()

    )

    for accion in sin_plan:

        try:

            limite = date.fromisoformat(accion.plazo_entrega_plan)

        except ValueError:

            continue

        if accion.ultima_notificacion_vencimiento == hoy_iso:

            continue  # ya se avisó hoy

        faltan = (limite - hoy).days

        if faltan < 0:

            _notificar_usuarios_de_area(

                db, accion.areas,

                f"Recordatorio: plazo vencido — {accion.numero}",

                f"La Acción de Mejora {accion.numero} debía entregarse el {accion.plazo_entrega_plan} y todavía no se completó el plan de acción.",

            )

            accion.ultima_notificacion_vencimiento = hoy_iso

            _agregar_historial(db, accion, f"Recordatorio automático de vencimiento enviado a {accion.area_responsable or 'Área Responsable'}")

        elif faltan == 0 or faltan in AVISO_DIAS_ANTES:

            cuando = "vence HOY" if faltan == 0 else f"vence en {faltan} día(s)"

            _notificar_usuarios_de_area(

                db, accion.areas,

                f"Tu formulario {cuando} — {accion.numero}",

                f"El plazo para llenar y enviar el formulario de la Acción de Mejora {accion.numero} {cuando} "

                f"({accion.plazo_entrega_plan}). Ingresa al sistema, completa el plan de acción y envíalo.",

            )

            accion.ultima_notificacion_vencimiento = hoy_iso

            _agregar_historial(db, accion, f"Aviso previo: el plazo para llenar el formulario {cuando}")



    # 2) Se cumplió el plazo de ejecución: ahora sí se piden las evidencias

    en_ejecucion = db.query(AccionMejora).filter(AccionMejora.estado == "En ejecución").all()

    for accion in en_ejecucion:

        try:

            fin = date.fromisoformat(accion.plazo_ejecucion or "")

        except ValueError:

            continue

        if fin > hoy:

            # Aún no se cumple: aviso previo para que vayan preparando las evidencias

            faltan = (fin - hoy).days

            if faltan in AVISO_DIAS_ANTES and accion.ultima_notificacion_evidencia != hoy_iso:

                _notificar_usuarios_de_area(

                    db, accion.areas,

                    f"Prepara tus evidencias (en {faltan} día(s)) — {accion.numero}",

                    f"El plazo de ejecución de la Acción de Mejora {accion.numero} se cumple el {accion.plazo_ejecucion} "

                    f"(en {faltan} día(s)). Ese día deberás subir al sistema la evidencia que indicaste en tu plan "

                    "y presionar «Enviar evidencia».",

                )

                accion.ultima_notificacion_evidencia = hoy_iso

                _agregar_historial(db, accion, f"Aviso previo: las evidencias se piden en {faltan} día(s)")

            continue

        # (desde aquí sigue igual que antes: accion.estado = "Pendiente de evidencia" ...) 

        accion.estado = "Pendiente de evidencia"

        accion.ultima_notificacion_evidencia = hoy_iso

        _agregar_historial(db, accion, f"Se cumplió el plazo de ejecución ({accion.plazo_ejecucion}): el área debe subir las evidencias")

        _notificar_usuarios_de_area(

            db, accion.areas,

            f"Sube tus evidencias — {accion.numero}",

            f"Se cumplió el plazo de ejecución ({accion.plazo_ejecucion}) de la Acción de Mejora {accion.numero}. "

            "Ingresa al sistema, sube la evidencia que indicaste en tu plan (PDF, Excel, fotos) y presiona «Enviar evidencia».",

        )



    # 3) Recordatorio de evidencia cada 3 días

    por_evidenciar = db.query(AccionMejora).filter(AccionMejora.estado == "Pendiente de evidencia").all()

    for accion in por_evidenciar:

        ultima = accion.ultima_notificacion_evidencia

        if ultima:

            try:

                if (hoy - date.fromisoformat(ultima)).days < 3:

                    continue

            except ValueError:

                pass

        _notificar_usuarios_de_area(

            db, accion.areas,

            f"Recordatorio: sube tu evidencia — {accion.numero}",

            f"La Acción de Mejora {accion.numero} sigue esperando tu evidencia (plazo de ejecución: {accion.plazo_ejecucion}). "

            "Ingresa al sistema, sube los archivos y presiona «Enviar evidencia».",

        )

        accion.ultima_notificacion_evidencia = hoy_iso

        _agregar_historial(db, accion, "Recordatorio automático: falta enviar la evidencia")

    db.commit()





# =========================================================================

# Solicitudes — módulo INDEPENDIENTE de Acciones de Mejora.

# =========================================================================



def _agregar_historial_solicitud(db: Session, solicitud: Solicitud, evento: str):

    db.add(HistorialSolicitud(solicitud_id=solicitud.id, fecha=_ahora(), evento=evento))





def _verificar_acceso_solicitud(usuario: Usuario, solicitud: Solicitud):

    if usuario.rol != "admin" and (solicitud.destinatario_correo or "").lower() != (usuario.correo or "").lower():

        raise HTTPException(status_code=403, detail="Esta solicitud no fue enviada a tu cuenta.")





@app.get("/api/solicitudes", response_model=List[SolicitudOut])

def listar_solicitudes(db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    _revisar_solicitudes_vencidas(db)

    query = db.query(Solicitud).options(joinedload(Solicitud.historial))

    if usuario.rol != "admin":

        query = query.filter(Solicitud.destinatario_correo == usuario.correo)

    return query.order_by(Solicitud.id.desc()).all()





@app.post("/api/solicitudes", response_model=SolicitudOut)

def crear_solicitud(data: SolicitudCreate, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    total = db.query(Solicitud).count()

    numero = f"SOL-{total + 1:04d}"

    solicitud = Solicitud(

        numero=numero,

        estado="Pendiente",

        fecha_envio=_ahora(),

        **data.dict(),

    )

    db.add(solicitud)

    db.commit()

    db.refresh(solicitud)

    _agregar_historial_solicitud(db, solicitud, f"Solicitud enviada a {data.destinatario_nombre or 'destinatario'}")

    db.commit()

    enviar_correo(

        data.destinatario_correo,

        f"Nueva solicitud — {numero}",

        f"Se te envió la solicitud {numero}: {data.asunto}",

    )

    db.refresh(solicitud)

    return solicitud





@app.put("/api/solicitudes/{solicitud_id}/responder", response_model=SolicitudOut)

def responder_solicitud(solicitud_id: int, data: SolicitudResponder, db: Session = Depends(get_db), usuario: Usuario = Depends(obtener_usuario_actual)):

    solicitud = db.query(Solicitud).filter(Solicitud.id == solicitud_id).first()

    if not solicitud:

        raise HTTPException(status_code=404, detail="Solicitud no encontrada")

    _verificar_acceso_solicitud(usuario, solicitud)



    ya_habia_respondido = solicitud.estado == "Respondida"



    for key, value in data.dict(exclude_unset=True).items():

        setattr(solicitud, key, value)

    solicitud.fecha_respuesta = _ahora()

    solicitud.estado = "Respondida"

    db.commit()

    if ya_habia_respondido:

        _agregar_historial_solicitud(db, solicitud, "Respuesta editada por el destinatario")

    else:

        _agregar_historial_solicitud(db, solicitud, "Respondida por el destinatario")

    db.commit()

    if ya_habia_respondido:

        _notificar_admins(db, f"Solicitud actualizada — {solicitud.numero}", f"El destinatario editó su respuesta a la solicitud {solicitud.numero}.")

    else:

        _notificar_admins(db, f"Solicitud respondida — {solicitud.numero}", f"La solicitud {solicitud.numero} fue respondida.")

    db.refresh(solicitud)

    return solicitud





@app.patch("/api/solicitudes/{solicitud_id}", response_model=SolicitudOut)

def editar_solicitud(solicitud_id: int, data: SolicitudEdicion, db: Session = Depends(get_db), usuario: Usuario = Depends(requerir_admin)):

    solicitud = db.query(Solicitud).filter(Solicitud.id == solicitud_id).first()

    if not solicitud:

        raise HTTPException(status_code=404, detail="Solicitud no encontrada")



    for key, value in data.dict(exclude_unset=True).items():

        setattr(solicitud, key, value)

    db.commit()

    _agregar_historial_solicitud(db, solicitud, "Editada manualmente por Sistemas de Gestión (calidad)")

    db.commit()

    db.refresh(solicitud)

    return solicitud





def _revisar_solicitudes_vencidas(db: Session):

    hoy = date.today().isoformat()

    vencidas = (

        db.query(Solicitud)

        .filter(

            Solicitud.estado == "Pendiente",

            Solicitud.plazo_respuesta.isnot(None),

            Solicitud.plazo_respuesta != "",

        )

        .all()

    )

    for s in vencidas:

        try:

            vencida = date.fromisoformat(s.plazo_respuesta) < date.today()

        except ValueError:

            continue

        if not vencida or s.ultima_notificacion_vencimiento == hoy:

            continue

        enviar_correo(

            s.destinatario_correo,

            f"Recordatorio: solicitud vencida — {s.numero}",

            f"La solicitud {s.numero} ('{s.asunto}') venció el {s.plazo_respuesta} y no se respondió.",

        )

        s.ultima_notificacion_vencimiento = hoy

        _agregar_historial_solicitud(db, s, f"Recordatorio automático de vencimiento enviado a {s.destinatario_nombre or 'destinatario'}")

    db.commit()